import hashlib
import json
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.catalogs.models import CollectionChannel, RevenueConcept
from apps.progression.models import ProgressionSettings
from .models import RevenueLine, RevenueLoad


MONEY_FIELDS = ("gross", "commission", "vat", "other_expenses")


def money(value, field):
    try:
        raw = str(value or 0).strip().replace("\u00a0", "").replace(" ", "")
        raw = "".join(character for character in raw if character.isdigit() or character in ",.-")
        if "," in raw and "." in raw:
            if raw.rfind(",") > raw.rfind("."):
                raw = raw.replace(".", "").replace(",", ".")
            else:
                raw = raw.replace(",", "")
        elif "," in raw:
            raw = raw.replace(".", "").replace(",", ".")
        number = Decimal(raw or "0")
    except (InvalidOperation, ValueError):
        raise ValueError(f"{field}: el valor no es numérico")
    if number < 0:
        raise ValueError(f"{field}: no puede ser negativo")
    # Excel puede entregar residuos binarios (por ejemplo 10.4999999998)
    # aunque la celda se vea con dos decimales. La fuente canónica se guarda
    # siempre redondeada a centavos.
    return number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def rows_from_workbook(upload):
    workbook = load_workbook(upload, read_only=True, data_only=True)
    if "Detalle" not in workbook.sheetnames:
        raise ValueError("El archivo no contiene la hoja Detalle")
    sheet = workbook["Detalle"]
    header_row = None
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row and str(row[0] or "").strip().upper() == "CODIGO_CONCEPTO":
            header_row = index
            break
    if not header_row:
        raise ValueError("No se encontró el encabezado CODIGO_CONCEPTO")
    rows = []
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        code = str(values[0] or "").strip().upper()
        if not code and rows:
            break
        if not code:
            continue
        rows.append({
            "code": code,
            "receipt_count": values[2] or 0,
            "gross": values[3] or 0,
            "commission": values[4] or 0,
            "vat": values[5] or 0,
            "other_expenses": values[6] or 0,
            "observations": values[7] or "",
        })
    return rows


@transaction.atomic
def create_load(payload, rows, source_name=""):
    required = ("revenue_date", "tgp_entry_date", "channel_id")
    for field in required:
        if not payload.get(field):
            raise ValueError(f"Falta el campo {field}")
    try:
        revenue_date = datetime.strptime(payload["revenue_date"], "%Y-%m-%d").date()
        tgp_date = datetime.strptime(payload["tgp_entry_date"], "%Y-%m-%d").date()
        cutoff_date = datetime.strptime(payload.get("cutoff_date") or payload["tgp_entry_date"], "%Y-%m-%d").date()
        cutoff_time = datetime.strptime(payload.get("cutoff_time") or "23:59", "%H:%M").time()
        channel = CollectionChannel.objects.get(pk=payload["channel_id"], active=True)
    except (ValueError, CollectionChannel.DoesNotExist):
        raise ValueError("Fecha o canal inválido")

    normalized = []
    seen = set()
    for raw in rows:
        code = str(raw.get("code", "")).strip().upper()
        if not code or code in seen:
            raise ValueError(f"Código vacío o repetido: {code or '(vacío)'}")
        seen.add(code)
        try:
            concept = RevenueConcept.objects.get(code=code, active=True)
            receipts = int(raw.get("receipt_count") or 0)
        except RevenueConcept.DoesNotExist:
            raise ValueError(f"Concepto desconocido: {code}")
        except (TypeError, ValueError):
            raise ValueError(f"Cantidad inválida para {code}")
        if receipts < 0:
            raise ValueError(f"Cantidad inválida para {code}")
        values = {field: money(raw.get(field), f"{code}.{field}") for field in MONEY_FIELDS}
        normalized.append({"concept": concept, "receipt_count": receipts, **values,
                           "observations": str(raw.get("observations") or "")[:500]})
    if not normalized:
        raise ValueError("La carga no contiene filas")

    digest_data = {
        "revenue_date": str(revenue_date), "tgp_entry_date": str(tgp_date),
        "channel": channel.pk,
        "rows": [{k: str(v.code if k == "concept" else v) for k, v in item.items()} for item in normalized],
    }
    digest = hashlib.sha256(json.dumps(digest_data, sort_keys=True).encode()).hexdigest()
    if RevenueLoad.objects.filter(content_hash=digest).exists():
        raise ValueError("Este contenido ya fue cargado exactamente")

    correction = None
    correction_id = payload.get("correction_of")
    if correction_id:
        correction = RevenueLoad.objects.filter(pk=correction_id, status="VALID").first()
        if not correction:
            raise ValueError("La carga a corregir no existe o ya no está vigente")

    configuration = ProgressionSettings.get_solo()
    delay = int(payload.get("publication_delay_minutes", settings.PUBLICATION_DELAY_MINUTES))
    cutoff_at = timezone.make_aware(datetime.combine(cutoff_date, cutoff_time))
    if cutoff_at <= timezone.now() + timedelta(minutes=delay):
        raise ValueError("La fecha y hora de corte deben ser posteriores al inicio de la progresión")

    load = RevenueLoad.objects.create(
        revenue_date=revenue_date, tgp_entry_date=tgp_date, channel=channel, cutoff_date=cutoff_date,
        cutoff_time=cutoff_time,
        publication_delay_minutes=delay,
        algorithm_snapshot=configuration.algorithm,
        interval_value_snapshot=configuration.interval_value,
        interval_unit_snapshot=configuration.interval_unit,
        max_variation_snapshot=configuration.max_variation_percent,
        randomness_intensity_snapshot=configuration.randomness_intensity_percent,
        smoothing_intervals_snapshot=configuration.smoothing_intervals,
        tv_animation_seconds_snapshot=configuration.tv_animation_seconds,
        gamma_snapshot=configuration.gamma,
        source_name=source_name, content_hash=digest, correction_of=correction,
        correction_note=str(payload.get("correction_note") or ""),
    )
    RevenueLine.objects.bulk_create([RevenueLine(load=load, **item) for item in normalized])
    if correction:
        correction.status = "REPLACED"
        correction.save(update_fields=["status"])
    return load


def serialize_load(load, include_lines=False):
    data = {
        "id": load.id, "revenue_date": load.revenue_date, "tgp_entry_date": load.tgp_entry_date,
        "channel": load.channel.name, "cutoff_date": load.cutoff_date, "cutoff_time": load.cutoff_time,
        "uploaded_at": load.uploaded_at, "publication_delay_minutes": load.publication_delay_minutes,
        "algorithm": load.algorithm_snapshot, "interval_value": load.interval_value_snapshot,
        "interval_unit": load.interval_unit_snapshot,
        "status": load.status, "source_name": load.source_name,
        "gross_total": f"{load.gross_total:.2f}", "commission_total": f"{load.commission_total:.2f}",
        "vat_total": f"{load.vat_total:.2f}", "net_total": f"{load.net_total:.2f}",
        "receipt_count": sum(line.receipt_count for line in load.lines.all()),
    }
    if include_lines:
        data["lines"] = [{
            "code": line.concept.code, "concept": line.concept.name,
            "receipt_count": line.receipt_count, "gross": f"{line.gross:.2f}",
            "commission": f"{line.commission:.2f}", "vat": f"{line.vat:.2f}",
            "other_expenses": f"{line.other_expenses:.2f}", "observations": line.observations,
        } for line in load.lines.select_related("concept")]
    return data
